"""Document ingestion and text extraction."""

from __future__ import annotations

from pathlib import Path

from trustdoc_ai.core.types import Document
from trustdoc_ai.core.utils import new_id


class DocIntelAgent:
    """Parse supported files into local text records.

    Optional dependencies are imported only when needed so the text demo stays
    lightweight and testable on machines without the full ML stack.
    """

    def parse(self, paths: list[str]) -> list[Document]:
        return [self._parse_one(Path(path)) for path in paths]

    def _parse_one(self, path: Path) -> Document:
        if not path.exists():
            raise FileNotFoundError(path)

        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8")
            page_count = 1
        elif suffix == ".pdf":
            text, page_count = self._parse_pdf(path)
        elif suffix == ".docx":
            text, page_count = self._parse_docx(path)
        elif suffix == ".xlsx":
            text, page_count = self._parse_xlsx(path)
        else:
            raise ValueError(f"Unsupported document type: {path.suffix}")

        return Document(
            doc_id=new_id("doc"),
            path=str(path.resolve()),
            name=path.name,
            text=text,
            page_count=page_count,
            layout={"source_type": suffix.lstrip(".") or "text"},
        )

    def _parse_pdf(self, path: Path) -> tuple[str, int]:
        try:
            import pdfplumber  # type: ignore
        except ImportError as exc:
            raise RuntimeError("Install pdfplumber to parse PDF files") from exc

        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        return "\n\n".join(pages), max(1, len(pages))

    def _parse_docx(self, path: Path) -> tuple[str, int]:
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise RuntimeError("Install python-docx to parse DOCX files") from exc

        document = docx.Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines), 1

    def _parse_xlsx(self, path: Path) -> tuple[str, int]:
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise RuntimeError("Install openpyxl to parse XLSX files") from exc

        workbook = openpyxl.load_workbook(str(path), data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines), len(workbook.worksheets)
